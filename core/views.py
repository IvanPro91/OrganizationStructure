# core/views.py
import json

from django.http import JsonResponse, HttpResponse
from django.views.generic import TemplateView
from django.contrib.auth.mixins import LoginRequiredMixin
from django.core.serializers.json import DjangoJSONEncoder
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, Border, Side, PatternFill
from openpyxl.utils import get_column_letter

from .models import ObjectType, OrganizationObject


class DashboardView(LoginRequiredMixin, TemplateView):
    """Главная страница с деревом объектов"""
    template_name = 'core/dashboard.html'

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)

        # Типы объектов
        context['object_types'] = ObjectType.objects.all()

        # Корневые объекты с предзагрузкой детей
        context['root_objects'] = OrganizationObject.objects.filter(
            parent=None
        ).select_related('object_type').prefetch_related('children__object_type')

        # Состояние развернутых узлов из сессии (можно сохранять)
        expanded = self.request.session.get('expanded_nodes', [])
        context['expanded_nodes'] = json.dumps(expanded)

        return context

    def post(self, request, *args, **kwargs):
        # Сохраняем состояние развернутых узлов
        if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
            data = json.loads(request.body)
            request.session['expanded_nodes'] = data.get('expanded', [])
            return JsonResponse({'status': 'ok'})

        return super().get(request, *args, **kwargs)

class TableView(LoginRequiredMixin, TemplateView):
    """Отображение объектов в виде таблиц по типам"""
    template_name = 'core/table_view.html'

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        types = ObjectType.objects.prefetch_related('fields').all()
        data = []
        for obj_type in types:
            objects = OrganizationObject.objects.filter(
                object_type=obj_type
            ).select_related('parent').prefetch_related('field_values__field')
            obj_list = []
            for obj in objects:
                values = {}
                for fv in obj.field_values.all():
                    values[fv.field.name] = fv.get_value()
                obj_list.append({
                    'id': obj.id,
                    'name': obj.name,
                    'status': obj.get_status_display(),
                    'parent': obj.parent.name if obj.parent else '',
                    'values': values,
                })
            data.append({
                'type': obj_type,
                'objects': obj_list,
                'fields': list(obj_type.fields.all()),
            })
        context['data'] = data
        return context


class ExportExcelView(LoginRequiredMixin, TemplateView):
    """Экспорт всех объектов в Excel с группировкой по типам"""

    def get(self, request, *args, **kwargs):
        wb = Workbook()
        # Удаляем стандартный лист, будем создавать для каждого типа отдельный
        wb.remove(wb.active)

        types = ObjectType.objects.prefetch_related('fields').all()
        for obj_type in types:
            # Пропускаем типы, у которых нет объектов
            objects = OrganizationObject.objects.filter(object_type=obj_type).select_related('parent').prefetch_related('field_values__field')
            if not objects.exists():
                continue

            ws = wb.create_sheet(title=obj_type.name[:31])  # ограничение длины имени листа

            # Заголовки
            headers = ['ID', 'Название', 'Статус', 'Родитель']
            fields = list(obj_type.fields.all())
            for f in fields:
                headers.append(f.name)

            # Стили
            header_font = Font(name='Times New Roman', size=12, bold=True)
            header_fill = PatternFill(start_color='FFC000', end_color='FFC000', fill_type='solid')
            header_alignment = Alignment(horizontal='center', vertical='center')
            thin_border = Border(
                left=Side(style='thin'),
                right=Side(style='thin'),
                top=Side(style='thin'),
                bottom=Side(style='thin')
            )

            # Записываем заголовки
            for col_num, header in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col_num, value=header)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = header_alignment
                cell.border = thin_border

            # Данные
            row_num = 2
            for obj in objects:
                # Основные поля
                for col_num, value in enumerate([obj.id, obj.name, obj.get_status_display(), obj.parent.name if obj.parent else ''], 1):
                    cell = ws.cell(row=row_num, column=col_num, value=value)
                    cell.border = thin_border
                    cell.alignment = Alignment(vertical='center')

                # Динамические поля
                values = {fv.field.name: fv.get_value() for fv in obj.field_values.all()}
                for col_num, field in enumerate(fields, 5):
                    val = values.get(field.name, '')
                    cell = ws.cell(row=row_num, column=col_num, value=val)
                    cell.border = thin_border
                    cell.alignment = Alignment(vertical='center')

                    # Специальная обработка булевых значений
                    if field.field_type == 'checkbox':
                        if val is True:
                            cell.fill = PatternFill(start_color='A9D08E', end_color='A9D08E', fill_type='solid')
                            cell.value = 'Да'
                        elif val is False:
                            cell.fill = PatternFill(start_color='FFE699', end_color='FFE699', fill_type='solid')
                            cell.value = 'Нет'
                        else:
                            cell.value = ''
                    elif field.field_type == 'datetime' and val:
                        # Форматирование даты
                        cell.value = val.strftime('%d.%m.%Y %H:%M')
                        cell.number_format = 'DD.MM.YYYY HH:MM'
                    # Для остальных типов оставляем как есть

                row_num += 1

            # Устанавливаем высоту строки заголовков (65 пунктов)
            ws.row_dimensions[1].height = 65

            # Автоширина колонок по содержимому
            for col in range(1, len(headers)+1):
                column_letter = get_column_letter(col)
                # Можно установить приблизительную ширину, но лучше вычислить максимальную длину
                max_length = 0
                for row in range(1, row_num):
                    cell_value = ws.cell(row=row, column=col).value
                    if cell_value:
                        max_length = max(max_length, len(str(cell_value)))
                # Учитываем заголовок
                header_len = len(headers[col-1])
                max_length = max(max_length, header_len)
                # Устанавливаем ширину с запасом
                ws.column_dimensions[column_letter].width = min(max_length + 2, 50)  # ограничим 50

        # Формируем ответ
        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename=objects_export.xlsx'
        wb.save(response)
        return response